import sys
import os
import time
import json
import re
import datetime
import subprocess
import urllib.request
import websocket
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

try:
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    sys.stderr.reconfigure(encoding='utf-8', errors='replace')
except Exception:
    pass

# Configure paths
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_DIR = os.path.abspath(os.path.join(SCRIPT_DIR, ".."))
DATA_JSON_PATH = os.path.join(PROJECT_DIR, "src", "data", "terrenos.json")
SYNC_STATE_PATH = os.path.join(PROJECT_DIR, "last_sync.json")
LOG_PATH = os.path.join(PROJECT_DIR, "sync_execution.log")

REGIONS = [
    {
        "city": "Chihuahua",
        "mkt_url": "https://www.facebook.com/marketplace/chihuahua/search/?query=terreno&sortBy=creation_time_descend",
        "search_url": "https://www.facebook.com/search/posts/?q=terreno%20venta%20chihuahua"
    },
    {
        "city": "Aldama",
        "mkt_url": "https://www.facebook.com/marketplace/search/?query=terreno%20aldama&sortBy=creation_time_descend",
        "search_url": "https://www.facebook.com/search/posts/?q=terreno%20aldama%20chihuahua%20venta"
    },
    {
        "city": "Delicias",
        "mkt_url": "https://www.facebook.com/marketplace/delicias/search/?query=terreno&sortBy=creation_time_descend",
        "search_url": "https://www.facebook.com/search/posts/?q=terreno%20delicias%20chihuahua%20venta"
    },
    {
        "city": "Meoqui",
        "mkt_url": "https://www.facebook.com/marketplace/search/?query=terreno%20meoqui&sortBy=creation_time_descend",
        "search_url": "https://www.facebook.com/search/posts/?q=terreno%20meoqui%20venta"
    },
    {
        "city": "Aquiles Serdán / Sta. Eulalia",
        "mkt_url": "https://www.facebook.com/marketplace/search/?query=terreno%20aquiles%20serdan&sortBy=creation_time_descend",
        "search_url": "https://www.facebook.com/search/posts/?q=terreno%20aquiles%20serdan%20santa%20eulalia"
    }
]

def log(msg):
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    formatted = f"[{now}] {msg}"
    try:
        print(formatted)
    except Exception:
        try:
            print(formatted.encode('ascii', errors='replace').decode('ascii'))
        except Exception:
            pass
    try:
        with open(LOG_PATH, "a", encoding="utf-8") as f:
            f.write(formatted + "\n")
    except Exception:
        pass

def parse_price(text):
    if not text:
        return 0
    t = text.lower().replace('\n', ' ')
    
    # Millions / MDP
    m_mdp = re.search(r'(\d+(?:[\.,]\d+)?)\s*(?:mdp|mill[oó]n(?:es)?)', t)
    if m_mdp:
        val = float(m_mdp.group(1).replace(',', '.'))
        return int(val * 1_000_000)
        
    # Thousands / Mil
    m_mil = re.search(r'(\d+(?:[\.,]\d+)?)\s*(?:mil\b|k\b)', t)
    if m_mil:
        val = float(m_mil.group(1).replace(',', '.'))
        return int(val * 1_000)

    # Standard currency $1,250,000
    m_std = re.search(r'\$\s*([\d\.,]+)', text)
    if m_std:
        clean = m_std.group(1).replace(',', '').strip()
        try:
            val = float(clean)
            if 10 <= val <= 999: # 380 -> 380,000
                val = val * 1000
            elif 1 <= val < 10:
                val = val * 1_000_000
            if val > 100:
                return int(val)
        except:
            pass

    # Numbers between 20,000 and 90,000,000
    m_digits = re.search(r'\b(\d{5,8})\b', text)
    if m_digits:
        val = int(m_digits.group(1))
        if 20_000 <= val <= 90_000_000:
            return val

    return 0

def parse_meters(text):
    if not text:
        return 0
    t = text.lower().replace('\n', ' ')

    # Hectáreas
    m_ha = re.search(r'(\d+(?:[\.,]\d+)?)\s*(?:ha\b|has\b|hect[aá]reas?)', t)
    if m_ha:
        val = float(m_ha.group(1).replace(',', '.'))
        if 0.1 <= val <= 5000:
            return round(val * 10_000, 2)

    # Dimensions AxB
    m_dim = re.search(r'(\d+(?:[\.,]\d+)?)\s*[xX*]\s*(\d+(?:[\.,]\d+)?)', t)
    if m_dim:
        try:
            w = float(m_dim.group(1).replace(',', '.'))
            l = float(m_dim.group(2).replace(',', '.'))
            if 3 <= w <= 500 and 3 <= l <= 1000:
                area = round(w * l, 2)
                if 20 <= area <= 500_000:
                    return area
        except:
            pass

    # Direct m2
    m_m2 = re.search(r'(\d+(?:[\.,\s]\d+)?)\s*(?:m2|m²|mts2|mts|metros\s*cuadrados)', t)
    if m_m2:
        clean = m_m2.group(1).replace(',', '').replace(' ', '').strip()
        try:
            val = float(clean)
            if 15 <= val <= 5_000_000:
                return round(val, 2)
        except:
            pass

    return 0

class ChromeTabCDP:
    def __init__(self, ws_url):
        self.ws_url = ws_url
        self.ws = websocket.create_connection(
            ws_url,
            timeout=25,
            suppress_origin=True,
            header={"User-Agent": "Mozilla/5.0"}
        )
        self.msg_id = 1

    def send_cmd(self, method, params=None):
        mid = self.msg_id
        self.msg_id += 1
        payload = {"id": mid, "method": method, "params": params or {}}
        self.ws.send(json.dumps(payload))
        while True:
            resp = json.loads(self.ws.recv())
            if resp.get("id") == mid:
                if "error" in resp:
                    raise Exception(resp["error"])
                return resp.get("result", {})

    def evaluate(self, expression):
        res = self.send_cmd("Runtime.evaluate", {
            "expression": expression,
            "returnByValue": True,
            "awaitPromise": True
        })
        result = res.get("result", {})
        return result.get("value")

    def navigate(self, url):
        self.send_cmd("Page.navigate", {"url": url})
        time.sleep(3.5)

    def close(self):
        try:
            self.ws.close()
        except Exception:
            pass

def try_get_cdp_tab():
    try:
        req = urllib.request.urlopen("http://127.0.0.1:9222/json", timeout=4)
        tabs = json.loads(req.read().decode())
        for t in tabs:
            if t.get("type") == "page" and "facebook.com" in t.get("url", ""):
                ws_url = t.get("webSocketDebuggerUrl")
                if ws_url:
                    return ChromeTabCDP(ws_url)
        req_new = urllib.request.Request("http://127.0.0.1:9222/json/new?https://www.facebook.com/marketplace", method="PUT")
        res = json.loads(urllib.request.urlopen(req_new).read().decode())
        return ChromeTabCDP(res.get("webSocketDebuggerUrl"))
    except Exception as e:
        log(f"CDP no disponible directamente ({e}), usando dataset base enriquecido")
        return None

def should_run_sync(force=False):
    if force:
        return True, "Forzado manualmente"
    if not os.path.exists(SYNC_STATE_PATH):
        return True, "Primera ejecución de sincronización"
    try:
        with open(SYNC_STATE_PATH, "r", encoding="utf-8") as f:
            state = json.load(f)
        last_str = state.get("last_sync_timestamp")
        if not last_str:
            return True, "Sin registro previo"
        last_time = datetime.datetime.fromisoformat(last_str)
        now = datetime.datetime.now()
        diff = now - last_time
        if diff.days >= 1:
            return True, f"Han pasado {diff.days} días desde la última sincronización"
        if diff.total_seconds() >= 21600: # 6 horas
            return True, f"Sincronización periódica (pasaron {int(diff.total_seconds()//3600)} horas)"
        return False, f"Ya se ejecutó hoy ({last_str})"
    except Exception as e:
        return True, f"Error al leer estado: {e}"

def run_sync(force=False):
    log("=======================================================")
    log("INICIANDO SUBPROCESO DE ACTUALIZACIÓN TERRENOS HUB")
    log("=======================================================")

    should_run, reason = should_run_sync(force=force)
    log(f"Motivo / Estado: {reason}")
    if not should_run:
        log("No es necesario ejecutar hoy. Finalizando.")
        return

    # Load current dataset
    current_items = []
    if os.path.exists(DATA_JSON_PATH):
        try:
            with open(DATA_JSON_PATH, "r", encoding="utf-8") as f:
                current_items = json.load(f)
            log(f"Dataset actual cargado: {len(current_items)} terrenos.")
        except Exception as e:
            log(f"Error al cargar dataset actual: {e}")

    seen_urls = {item.get("url") for item in current_items if item.get("url")}
    new_extracted = []

    # Scraping pass if Chrome is open
    cdp = try_get_cdp_tab()
    if cdp:
        try:
            log("Escaneando publicaciones recientes en Facebook...")
            for reg in REGIONS:
                city = reg["city"]
                cdp.navigate(reg["mkt_url"])
                time.sleep(2.5)
                for _ in range(5):
                    cdp.evaluate("window.scrollBy(0, 1500);")
                    time.sleep(1.0)

                items = cdp.evaluate("""() => {
                    const links = Array.from(document.querySelectorAll('a[href*="/marketplace/item/"]'));
                    return links.map(a => ({
                        href: a.href,
                        text: a.innerText || ''
                    }));
                }""") or []

                for it in items:
                    raw_href = it.get("href", "").split("?")[0]
                    if not raw_href or raw_href in seen_urls:
                        continue
                    seen_urls.add(raw_href)
                    txt = it.get("text", "")
                    lines = [l.strip() for l in txt.splitlines() if l.strip()]
                    title = lines[0] if lines else "Terreno en venta"
                    p_val = parse_price(txt)
                    m_val = parse_meters(txt)
                    
                    new_extracted.append({
                        "city": city,
                        "title": title[:100],
                        "meters": m_val,
                        "price": p_val,
                        "location_card": f"{city} / Zona",
                        "date": datetime.datetime.now().strftime("%Y-%m-%d"),
                        "url": raw_href,
                        "description": txt[:1000],
                        "source": "Facebook Marketplace",
                        "coordinates": None,
                        "status": "Disponible",
                        "notes": ""
                    })
            log(f"Nuevos terrenos capturados en vivo: {len(new_extracted)}")
        except Exception as e:
            log(f"Aviso durante escaneo en vivo: {e}")
        finally:
            cdp.close()

    # Merge new items
    all_items = new_extracted + current_items

    # Clean, calculate $/m2 and sort by size (meters desc)
    for it in all_items:
        p = it.get("price", 0)
        m = it.get("meters", 0)
        if p > 0 and m > 0:
            it["price_per_m2"] = round(p / m, 2)
        else:
            it["price_per_m2"] = 0

    all_items.sort(key=lambda x: (x.get("meters", 0), x.get("price", 0)), reverse=True)

    for idx, it in enumerate(all_items):
        it["id"] = f"TER-{idx+1:03d}"

    # Save JSON to web app
    os.makedirs(os.path.dirname(DATA_JSON_PATH), exist_ok=True)
    with open(DATA_JSON_PATH, "w", encoding="utf-8") as f:
        json.dump(all_items, f, ensure_ascii=False, indent=2)
    log(f"✅ Archivo JSON actualizado con {len(all_items)} registros.")

    # Generate Excel on Desktop and project
    export_excel(all_items)

    # Git Commit & Push (Auto Vercel Deploy)
    push_to_git()

    # Save last sync state
    with open(SYNC_STATE_PATH, "w", encoding="utf-8") as f:
        json.dump({
            "last_sync_timestamp": datetime.datetime.now().isoformat(),
            "total_terrenos": len(all_items),
            "new_added": len(new_extracted),
            "status": "success"
        }, f, indent=2)
    log("🎉 Subproceso completado exitosamente.")

def export_excel(terrenos):
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Terrenos por Superficie"

        header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="E5A93C")
        border_thin = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        headers = [
            "ID", "Superficie (m²)", "Precio ($ MXN)", "Precio / m² ($)", "Ciudad / Municipio",
            "Título / Propiedad", "Zona / Ubicación", "Fecha Publicación", "Enlace Facebook", "Descripción"
        ]
        ws.append(headers)

        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for r_idx, t in enumerate(terrenos, start=2):
            ws.append([
                t["id"],
                t["meters"] if t["meters"] > 0 else "N/D",
                f"${t['price']:,.0f}" if t['price'] > 0 else "A cotizar",
                f"${t['price_per_m2']:,.2f}" if t['price_per_m2'] > 0 else "-",
                t["city"],
                t["title"],
                t.get("location_card", ""),
                t.get("date", ""),
                t.get("url", ""),
                t.get("description", "").replace("\n", " ")[:1000]
            ])
            fill = PatternFill(start_color="F8FAFC" if r_idx % 2 == 0 else "FFFFFF", fill_type="solid")
            for col in range(1, len(headers) + 1):
                cell = ws.cell(row=r_idx, column=col)
                cell.fill = fill
                cell.border = border_thin
                cell.alignment = Alignment(vertical="center")
                if col == 3:
                    cell.font = Font(name="Segoe UI", size=10, bold=True, color="D97706")
                elif col == 4:
                    cell.font = Font(name="Segoe UI", size=10, bold=True, color="059669")
                elif col == 9:
                    cell.font = Font(name="Segoe UI", size=9, color="2563EB", underline="single")

        ws.column_dimensions['A'].width = 10
        ws.column_dimensions['B'].width = 18
        ws.column_dimensions['C'].width = 18
        ws.column_dimensions['D'].width = 16
        ws.column_dimensions['E'].width = 20
        ws.column_dimensions['F'].width = 35
        ws.column_dimensions['G'].width = 24
        ws.column_dimensions['H'].width = 22
        ws.column_dimensions['I'].width = 36
        ws.column_dimensions['J'].width = 60

        out_paths = [
            r"C:\Users\webut\Desktop\Terrenos_Chihuahua_Delicias_Aldama.xlsx",
            r"C:\Users\webut\Desktop\04_Bienes_Raices\Terrenos_Chihuahua_Delicias_Aldama.xlsx"
        ]
        for p in out_paths:
            os.makedirs(os.path.dirname(p), exist_ok=True)
            wb.save(p)
            log(f"📁 Excel guardado en: {p}")
    except Exception as e:
        log(f"Error al exportar Excel: {e}")

def push_to_git():
    try:
        log("Sincronizando con GitHub para auto-despliegue en Vercel...")
        subprocess.run(["git", "add", "."], cwd=PROJECT_DIR, check=True)
        
        diff = subprocess.run(["git", "diff", "--cached", "--quiet"], cwd=PROJECT_DIR)
        if diff.returncode == 0:
            log("No hay cambios pendientes de datos para enviar a Git.")
            return

        date_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
        subprocess.run(["git", "commit", "-m", f"auto: daily sync {date_str} - Vercel auto-deploy"], cwd=PROJECT_DIR, check=True)
        subprocess.run(["git", "push", "origin", "main"], cwd=PROJECT_DIR, check=True)
        log("🚀 Push exitoso a GitHub -> Vercel se está actualizando automáticamente en vivo.")
    except Exception as e:
        log(f"Aviso al hacer push a Git: {e}")

if __name__ == "__main__":
    force_run = "--force" in sys.argv
    run_sync(force=force_run)
